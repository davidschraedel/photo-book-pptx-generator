# from pptx import Presentation

# prs = Presentation("book_template.pptx")
# # Update this to match how you grab your layout
# layout = prs.slide_layouts[7] 

# print("Placeholder Types:", [ph.placeholder_format.type for ph in layout.placeholders])



from pptx import Presentation

def get_layout_by_name(prs, name):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise ValueError(f"Layout '{name}' not found")

prs = Presentation("book_template.pptx")
grid_layout = get_layout_by_name(prs, "Custom Layout")

print("All elements on 'Custom Layout':")
for shape in grid_layout.shapes:
    print(f"- Name: '{shape.name}' | Is Placeholder: {shape.is_placeholder} | Type: {shape.shape_type}")