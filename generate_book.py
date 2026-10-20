from pathlib import Path
from PIL import Image
from pptx import Presentation
import io
import rawpy
import shutil
from datetime import datetime

# Ensure all image plugins are loaded before grabbing valid extensions
Image.init()
VALID_EXTENSIONS = {ext.lower() for ext in Image.registered_extensions().keys()} | {".cr2"}

SOURCE_TEMPLATE = Path("book_template.pptx")
PHOTOS_DIR = Path("photos/current")
PHOTOS_USED_DIR = Path("photos/used")
OUTPUT_FILE = Path(f"output/book_{PHOTOS_DIR.name}.pptx")

MAX_PAGES = 70
MAX_PHOTOS = MAX_PAGES * 4  # 280

def make_used_run_directory() -> Path:
    """One folder per book run so filenames stay unchanged (no _2/_3 renaming)."""
    PHOTOS_USED_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    run_dir = PHOTOS_USED_DIR / f"run_{stamp}"
    n = 2
    while run_dir.exists():
        run_dir = PHOTOS_USED_DIR / f"run_{stamp}_{n}"
        n += 1
    run_dir.mkdir()
    return run_dir

def used_photo_dest(run_dir: Path, filename: str) -> Path:
    dest = run_dir / filename
    if not dest.exists():
        return dest
    path = Path(filename)
    while True:
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S.%f")
        dest = run_dir / f"{path.stem}_{stamp}{path.suffix}"
        if not dest.exists():
            return dest

def get_sorted_photos(folder: Path):
    if not folder.exists():
        raise FileNotFoundError(f"Directory not found: {folder}")
    
    photos = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in VALID_EXTENSIONS]
    # Sort by filename, ascending — matches how photos are already organized
    return sorted(photos, key=lambda p: p.name)

def fit_and_center(img_source, box_left, box_top, box_width, box_height):
    if isinstance(img_source, Path):
        with Image.open(img_source) as img:
            img_w, img_h = img.size
    else:
        img_w, img_h = img_source.size
    
    scale = min(box_width / img_w, box_height / img_h)
    new_w, new_h = int(img_w * scale), int(img_h * scale)
    left = box_left + (box_width - new_w) // 2
    top = box_top + (box_height - new_h) // 2
    
    return left, top, new_w, new_h

def get_layout_by_name(prs, name):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise ValueError(f"Layout '{name}' not found in template")

def main():
    prs = Presentation(SOURCE_TEMPLATE)
    grid_layout = get_layout_by_name(prs, "Custom Layout")

    print(f"Scanning '{PHOTOS_DIR}' for photos...")
    photos = get_sorted_photos(PHOTOS_DIR)[:MAX_PHOTOS]
    
    if not photos:
        print("No valid photos found in the target directory. Exiting.")
        return

    groups = [photos[i:i+4] for i in range(0, len(photos), 4)]
    
    print(f"Found {len(photos)} photos. Preparing to generate {len(groups)} slides.")
    print("-" * 40)

    used_paths: list[Path] = []

    # Use enumerate to track the current slide number (starting at 1)
    for slide_num, group in enumerate(groups, start=1):
        print(f"Building Slide {slide_num}/{len(groups)}...")
        slide = prs.slides.add_slide(grid_layout)
        
        # Grab all placeholders, bypassing the broken type filters
        placeholders = list(slide.placeholders)
        
        # Sort placeholders top-to-bottom, then left-to-right
        placeholders.sort(key=lambda p: (p.top, p.left))

        for ph, photo_path in zip(placeholders, group):
            print(f"  -> Inserting: {photo_path.name}")
            
            box_left, box_top = ph.left, ph.top
            box_width, box_height = ph.width, ph.height
            
            # --- CR2 / RAW HANDLING ---
            if photo_path.suffix.lower() == ".cr2":
                with rawpy.imread(str(photo_path)) as raw:
                    rgb = raw.postprocess(use_camera_wb=True)
                
                pil_img = Image.fromarray(rgb)
                left, top, w, h = fit_and_center(pil_img, box_left, box_top, box_width, box_height)
                
                img_stream = io.BytesIO()
                pil_img.convert("RGB").save(img_stream, format="JPEG")
                img_stream.seek(0)
                slide.shapes.add_picture(img_stream, left, top, width=w, height=h)
                
            # --- STANDARD FORMAT HANDLING ---
            else:
                left, top, w, h = fit_and_center(photo_path, box_left, box_top, box_width, box_height)
                
                try:
                    slide.shapes.add_picture(str(photo_path), left, top, width=w, height=h)
                except ValueError as e:
                    if "unsupported image format" in str(e):
                        with Image.open(photo_path) as img:
                            img_stream = io.BytesIO()
                            img.convert("RGB").save(img_stream, format="JPEG")
                            img_stream.seek(0)
                            slide.shapes.add_picture(img_stream, left, top, width=w, height=h)
                    else:
                        raise

            used_paths.append(photo_path)

            # Remove empty placeholder outline
            ph._element.getparent().remove(ph._element)

    print("-" * 40)
    print("Saving presentation. This may take a moment depending on file size...")
    prs.save(OUTPUT_FILE)

    if used_paths:
        used_run_dir = make_used_run_directory()
        moves = [(src, used_photo_dest(used_run_dir, src.name)) for src in used_paths]
        print(f"Moving {len(moves)} used photo(s) to '{used_run_dir}'...")
        moved = 0
        for src, dest in moves:
            try:
                shutil.move(str(src), str(dest))
            except OSError as e:
                print("-" * 40)
                print("ERROR: Photo move failed partway through.")
                print(f"  Book already saved: {OUTPUT_FILE}")
                print(f"  Run folder: {used_run_dir}")
                print(f"  Moved {moved} of {len(moves)} photo(s).")
                print(f"  Failed on: {src.name} -> {dest}")
                print(f"  Reason: {e}")
                raise SystemExit(1)
            if dest.name == src.name:
                print(f"  -> Moved: {src.name}")
            else:
                print(f"  -> Moved: {src.name} -> {dest.name} (name conflict)")
            moved += 1

    print(f"Done: {len(used_paths)} photos across {len(groups)} pages -> {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
